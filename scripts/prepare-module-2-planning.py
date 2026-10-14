"""Author shared bilingual planning resources using the retained document styles."""
import argparse
import importlib.util
import json
import sys
from pathlib import Path

from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location("planning_styles", ROOT / "scripts/prepare-module-4-lesson-plans.py")
styles = importlib.util.module_from_spec(spec)
sys.dont_write_bytecode = True
spec.loader.exec_module(styles)
p, pair, field, heading, link = styles.paragraph, styles.bilingual, styles.blank, styles.heading, styles.link
OUT = ROOT / "public/assets/lesson-plans/module-2"
GUIDANCE = "https://learning.cambridgeinternational.org/classroom/course/view.php?id=4872"
TRANSCRIPT = "https://learning.cambridgeinternational.org/classroom/pluginfile.php/240077/mod_label/intro/Lesson%20planning%20transcript.pdf"
BOOK = "https://doi.org/10.1017/9781009024792.007"
MINUTES = [5, 5, 7, 10, 8, 5, 5]
STAGES = [
    ("Retrieve and orient", "Ôn tập và định hướng"),
    ("Prepare essential language", "Chuẩn bị ngôn ngữ cốt lõi"),
    ("Model the new thinking", "Làm mẫu cách tư duy mới"),
    ("Practise with support", "Luyện tập có hỗ trợ"),
    ("Apply independently", "Vận dụng độc lập"),
    ("Review and improve", "Rà soát và cải thiện"),
    ("Check and close", "Kiểm tra và kết thúc"),
]


def table(doc, headers, rows, widths, font_size=10.5):
    result = styles.table(doc, headers, rows, widths, font_size)
    gap = doc.add_paragraph()
    gap.paragraph_format.line_spacing = Pt(6)
    gap.paragraph_format.space_after = Pt(0)
    return result


def new(kind, title, vi):
    doc = styles.new_document({"name": "All subjects", "topic": [title]}, kind)
    doc.core_properties.title = title
    doc.core_properties.subject = "Shared 45 minute English Vietnamese EMI lesson planning"
    doc.styles["Normal"].paragraph_format.line_spacing = 1.05
    doc.styles["Normal"].paragraph_format.space_after = Pt(5)
    doc.styles["Subtitle"].paragraph_format.space_after = Pt(4)
    doc.styles["Reference"].font.size = Pt(10)
    p(doc, title, "Title")
    p(doc, vi, "Subtitle")
    pair(doc, ["For Maths, Physics, Chemistry, Biology, Engineering and IT. One lesson lasts 45 minutes.", "Dùng cho Toán học, Vật lí, Hóa học, Sinh học, Kĩ thuật và Công nghệ thông tin. Một tiết học kéo dài 45 phút."])
    p(doc, "Independent Cambridge-aligned resource / Tài liệu biên soạn độc lập theo định hướng Cambridge", "Reference")
    return doc


def sources(doc):
    p(doc, "Sources and scope", "Heading 2")
    p(doc, "Nguồn tham chiếu và phạm vi", "Subtitle")
    pair(doc, ["Adapted from Cambridge online planning guidance and Kay Bentley, The TKT Course CLIL Module, Unit 6, pp. 30-36, Cambridge University Press, 2010. The public chapter summary was consulted. Bilingual prompts, examples and 45-minute timings are original adaptations, not Cambridge requirements.", "Tham khảo hướng dẫn trực tuyến Cambridge và sách The TKT Course CLIL Module của Kay Bentley, Bài 6, trang 30-36, Cambridge University Press, 2010, qua tóm tắt chương công khai. Gợi ý song ngữ, ví dụ và lịch 45 phút được soạn độc lập, không phải yêu cầu Cambridge."]).style = "Reference"
    link(doc, "Cambridge lesson planning / Hướng dẫn lập kế hoạch bài dạy", GUIDANCE)
    link(doc, "Cambridge video transcript / Bản chép lời video hướng dẫn", TRANSCRIPT)
    link(doc, "The TKT Course CLIL Module Unit 6 / Bài 6 về lập kế hoạch CLIL", BOOK)
    pair(doc, ["Not issued or endorsed by Cambridge. Follow the school's format and applicable curriculum. Reference a lawfully obtained book; do not reproduce it wholesale. Sources checked 10 October 2026.", "Không do Cambridge phát hành hoặc chứng nhận. Theo mẫu của trường và chương trình áp dụng. Tham chiếu sách dùng hợp pháp, không sao chép toàn bộ. Kiểm tra nguồn ngày 10 tháng 10 năm 2026."]).style = "Reference"


def template():
    d = new("template", "Shared Lesson Plan Format", "Mẫu kế hoạch bài dạy dùng chung")
    pair(d, ["Fill the Word format and expand the spaces as needed. Connect outcomes, evidence and activities. The PDF is for reading or printing, not an interactive form.", "Điền mẫu Word và mở rộng chỗ điền nếu cần. Liên kết mục tiêu, bằng chứng và hoạt động. PDF để đọc hoặc in, không phải biểu mẫu tương tác."])
    p(d, "1 Lesson and book information", "Heading 1")
    p(d, "Thông tin bài dạy và sách", "Subtitle")
    for en, vi in [("School and teacher", "Trường và giáo viên"), ("Subject and title", "Môn và tên bài"), ("Date class age group", "Ngày lớp độ tuổi"), ("Present and absent", "Có mặt và vắng")]:
        p(d, f"{en} / {vi}: ...............................................")
    field(d, "Book title author edition chapter pages", "Tên sách tác giả phiên bản chương trang")
    field(d, "Curriculum objective and unit link", "Yêu cầu chương trình và liên kết chủ đề")
    p(d, "Duration / Thời lượng: 45 minutes / phút", bold=True)
    p(d, "2 Outcomes and success criteria", "Heading 1")
    p(d, "Mục tiêu và tiêu chí thành công", "Subtitle")
    table(d, ["Learning outcome\nMục tiêu học tập", "Success criterion\nTiêu chí thành công", "Evidence\nBằng chứng"], [
        ["Content 1 / Nội dung 1\n...............................", "Observable feature / Dấu hiệu quan sát\n...............................", "Task or response / Bài làm\n........................"],
        ["Content 2 if needed / Nội dung 2 nếu cần\n...............................", "Observable feature / Dấu hiệu quan sát\n...............................", "Task or response / Bài làm\n........................"],
        ["Language / Ngôn ngữ\n...............................", "Accurate use / Sử dụng chính xác\n...............................", "Spoken or written / Nói hoặc viết\n........................"],
    ], [2.6, 2.65, 1.95], 10.5)

    heading(d, "3 Starting points and language", "Kiến thức nền và ngôn ngữ", new_page=True)
    field(d, "Prior knowledge and opening diagnostic question", "Kiến thức nền và câu hỏi kiểm tra đầu giờ", 2)
    field(d, "Likely misconception and how to check it", "Hiểu lầm dự kiến và cách kiểm tra", 2)
    field(d, "Class language profile and access needs", "Đặc điểm ngôn ngữ và nhu cầu tiếp cận của lớp")
    p(d, "Essential vocabulary", "Heading 2")
    p(d, "Từ vựng cốt lõi", "Subtitle")
    table(d, ["English term\nThuật ngữ tiếng Anh", "Vietnamese meaning\nNghĩa tiếng Việt", "Visual example or pronunciation cue\nHình ví dụ hoặc gợi ý phát âm"], [["........................", "........................", "................................"] for _ in range(5)], [2.1, 2.1, 3.0], 11)
    field(d, "Language function and sentence frame", "Chức năng ngôn ngữ và khung câu", 2)
    field(d, "Visual gesture or worked example", "Hình ảnh cử chỉ hoặc ví dụ mẫu")
    field(d, "When Vietnamese support is needed and the English return task", "Khi cần hỗ trợ tiếng Việt và nhiệm vụ chuyển lại tiếng Anh", 2)
    field(d, "Resources quantities preparation and textbook task references", "Học liệu số lượng chuẩn bị và vị trí bài tập trong sách", 2)

    heading(d, "4 The 45 minute lesson sequence", "Tiến trình tiết học 45 phút", new_page=True)
    pair(d, ["Suggested pacing: beginning 5 minutes, middle 35 minutes, end 5 minutes. Adapt the activity stages to the subject while keeping the total at 45 minutes. Include transitions, setup and any cleanup within these times.", "Phân bổ gợi ý: mở đầu 5 phút, phần chính 35 phút, kết thúc 5 phút. Điều chỉnh các hoạt động theo môn nhưng giữ tổng 45 phút. Tính cả chuyển hoạt động, chuẩn bị và thu dọn vào các khoảng này."])
    rows, start = [], 0
    for minutes, (en, vi) in zip(MINUTES, STAGES):
        rows.append([f"{start}-{start + minutes}\n{minutes} min / phút", f"{en}\n{vi}\nTeacher / GV: ..................", "Task and grouping\nNhiệm vụ và hình thức\n................................", "Question and response\nCâu hỏi và phản hồi\n............................"])
        start += minutes
    assert start == 45
    table(d, ["Time\nThời gian", "Stage and teacher action\nGiai đoạn và việc của GV", "Learner action\nViệc của học sinh", "Evidence and next action\nBằng chứng và xử lí"], rows, [.8, 2.5, 2.1, 1.8], 10.5)
    field(d, "One key question with expected answer", "Một câu hỏi trọng tâm và đáp án dự kiến", 2)
    field(d, "If learners need longer which optional task can be shortened", "Nếu học sinh cần thêm thời gian thì rút gọn hoạt động nào", 2)

    heading(d, "5 Support assessment and follow up", "Hỗ trợ đánh giá và hoạt động tiếp nối", new_page=True)
    field(d, "Support for content or English access", "Hỗ trợ để tiếp cận kiến thức hoặc tiếng Anh", 2)
    field(d, "Extension that deepens reasoning", "Nhiệm vụ mở rộng để đào sâu lập luận", 2)
    field(d, "Safety accessibility or digital privacy and offline backup", "An toàn khả năng tiếp cận hoặc riêng tư số và phương án không có mạng", 2)
    field(d, "Mid lesson check expected response and reteaching trigger", "Kiểm tra giữa giờ đáp án và dấu hiệu cần dạy lại", 2)
    field(d, "Individual exit task answer guidance and success criteria", "Nhiệm vụ cuối giờ cá nhân hướng dẫn đáp án và tiêu chí", 2)
    field(d, "Feedback and how learners act on it", "Phản hồi và cách học sinh sử dụng phản hồi")
    field(d, "Homework if appropriate and link to the next lesson", "Bài về nhà nếu phù hợp và liên kết bài sau", 2)

    heading(d, "6 Reflection and evaluation", "Phản tư và đánh giá sau giờ dạy", new_page=True)
    pair(d, ["Complete after teaching with actual evidence. Do not pre-fill attainment or assume that an activity proves learning.", "Điền sau khi dạy bằng bằng chứng thực tế. Không ghi trước kết quả hoặc mặc định làm hoạt động là đã học được."])
    field(d, "Outcomes achieved and evidence collected", "Mục tiêu đạt và bằng chứng thu thập", 2)
    field(d, "Learners who need a next step and the support planned", "Học sinh cần bước tiếp theo và hỗ trợ dự kiến", 2)
    field(d, "Language support and timing that worked or need adjustment", "Hỗ trợ ngôn ngữ và thời gian hiệu quả hoặc cần điều chỉnh", 2)
    field(d, "One change for the next lesson with a reason", "Một điều chỉnh cho bài sau kèm lý do", 2)
    sources(d)
    return d


def guide():
    content = json.loads((ROOT / "scripts/data/module-2-step-guide.json").read_text())
    assert content["lessonMinutes"] == sum(MINUTES) == 45
    d = new("guide", "Step by Step Lesson Planning Guide", "Hướng dẫn từng bước xây dựng kế hoạch bài dạy")
    for index, page in enumerate(content["pages"]):
        if index:
            heading(d, page["heading"], page["vi"], new_page=True)
        for block in page["blocks"]:
            kind = block["type"]
            if kind == "pair":
                paragraph = pair(d, block["text"])
                paragraph.paragraph_format.keep_together = True
                if page.get("reference"):
                    paragraph.style = "Reference"
            elif kind == "heading":
                p(d, block["en"], "Heading 2")
                p(d, block["vi"], "Subtitle")
            elif kind == "actions":
                for number, item in enumerate(block["items"], 1):
                    pair(d, [f"{number}. {item[0]}", f"{number}. {item[1]}"])
            elif kind == "table":
                table(d, block["headers"], block["rows"], block["widths"], 10.5)
            elif kind == "subject":
                p(d, block["en"], "Heading 2")
                p(d, block["vi"], "Subtitle")
                table(d, ["Planning choice\nLựa chọn khi soạn", "Subject example\nVí dụ môn học"], block["rows"], [2.0, 5.2], 10.5)
            elif kind == "field":
                field(d, block["en"], block["vi"], block["lines"])
            elif kind == "link":
                link(d, block["label"], block["url"])
            else:
                raise ValueError(f"Unknown guide block: {kind}")
        # Do not let a spacer after a full-page table create an empty page.
        if d.paragraphs and not d.paragraphs[-1].text:
            spacer = d.paragraphs[-1]._element
            spacer.getparent().remove(spacer)
    return d


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--kind", choices=["all", "format", "guide"], default="all")
    args = parser.parse_args()
    OUT.mkdir(parents=True, exist_ok=True)
    for kind, stem, build in [("format", "shared-lesson-plan-format-en-vi", template), ("guide", "step-by-step-lesson-planning-guide-en-vi", guide)]:
        if args.kind not in {"all", kind}:
            continue
        path = OUT / f"{stem}.docx"
        build().save(path)
        print(path)


if __name__ == "__main__":
    main()
