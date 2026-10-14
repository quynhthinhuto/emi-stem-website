"""Build independently authored bilingual Cambridge-aligned planning resources."""
import argparse
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DATA = json.loads((ROOT / "scripts/data/module-4-lesson-plans.json").read_text())
OUT = ROOT / "public/assets/lesson-plans/module-4"
COURSE = DATA["guidance"]
LESSON_MINUTES = DATA["lessonMinutes"]
SITE = "https://emi-stem-vietnam-training-20260913.quynhthinhuto.chatgpt.site"


def paragraph(doc, text="", style=None, bold=None):
    if style in {"Title", "Subtitle", "Heading 1", "Heading 2"}:
        text = text.replace(" / ", "   ").replace("  |  ", "   ")
    p = doc.add_paragraph(style=style)
    p.add_run(text).bold = bold
    return p


def bilingual(doc, pair, label=None):
    if label:
        paragraph(doc, label, "Heading 2")
    p = doc.add_paragraph()
    p.add_run(pair[0])
    r = p.add_run("\n" + pair[1])
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "vi-VN")
    r._r.get_or_add_rPr().append(lang)
    return p


def blank(doc, en, vi, lines=1):
    p = doc.add_paragraph()
    p.add_run(f"{en} / {vi}").bold = True
    for _ in range(lines):
        p.add_run("\n........................................................................................")


def heading(doc, en, vi, new_page=False):
    p = paragraph(doc, en, "Heading 1")
    if new_page:
        p.paragraph_format.page_break_before = True
    p = paragraph(doc, vi, "Subtitle")
    p.paragraph_format.space_after = Pt(9)


def table(doc, headers, rows, widths, font_size=10.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    props = t._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        el = OxmlElement(f"w:{edge}")
        for attr, value in {"val": "single", "sz": "4", "color": "D9D9D9"}.items():
            el.set(qn(f"w:{attr}"), value)
        borders.append(el)
    props.append(borders)
    margins = OxmlElement("w:tblCellMar")
    for side in ["top", "bottom", "left", "right"]:
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), "80")
        el.set(qn("w:type"), "dxa")
        margins.append(el)
    props.append(margins)
    repeat = OxmlElement("w:tblHeader")
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row_index, values in enumerate([headers, *rows]):
        cells = t.rows[0].cells if row_index == 0 else t.add_row().cells
        no_split = OxmlElement("w:cantSplit")
        cells[0]._tc.getparent().get_or_add_trPr().append(no_split)
        for index, (cell, value) in enumerate(zip(cells, values)):
            cell.width = Inches(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            p.paragraph_format.keep_with_next = False
            r = p.add_run(value)
            r.font.size = Pt(font_size)
            if row_index == 0:
                r.bold = True
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "F2F2F2")
                cell._tc.get_or_add_tcPr().append(shade)
            if index == 0 and len(headers) == 3:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    return t


def link(doc, label, url):
    p = doc.add_paragraph()
    rel = doc.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rel)
    r = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "006D77")
    props.append(color)
    r.append(props)
    text = OxmlElement("w:t")
    text.text = label
    r.append(text)
    h.append(r)
    p._p.append(h)
    p.style = "Reference"


def new_document(s, kind):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin, sec.bottom_margin = Inches(.6), Inches(.6)
    sec.left_margin, sec.right_margin = Inches(.65), Inches(.65)
    sec.header_distance, sec.footer_distance = Inches(.25), Inches(.25)
    styles = doc.styles
    for style in styles:
        for border in list(style.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2"]:
        st = styles[name]
        st.font.name = "Noto Sans"
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.paragraph_format.line_spacing = 1.08
    styles["Normal"].font.size = Pt(11)
    styles["Normal"].paragraph_format.space_after = Pt(7)
    styles["Title"].font.size = Pt(23)
    styles["Title"].font.bold = True
    styles["Title"].paragraph_format.space_after = Pt(4)
    styles["Subtitle"].font.size = Pt(12)
    styles["Subtitle"].paragraph_format.space_after = Pt(7)
    for name, size in [("Heading 1", 15), ("Heading 2", 11.5)]:
        styles[name].font.size = Pt(size)
        styles[name].font.bold = True
        styles[name].paragraph_format.space_before = Pt(9)
        styles[name].paragraph_format.space_after = Pt(5)
    styles.add_style("Reference", 1)
    styles["Reference"].base_style = styles["Normal"]
    styles["Reference"].font.size = Pt(9)
    styles["Reference"].paragraph_format.space_after = Pt(4)
    doc.core_properties.title = f"{s['name']} bilingual lesson plan {kind}"
    doc.core_properties.subject = s["topic"][0]
    doc.core_properties.author = "EMI STEM Teacher Training"
    doc.core_properties.keywords = "English Vietnamese; bilingual; independently authored; lesson planning"
    footer = sec.footer.paragraphs[0]
    footer.paragraph_format.space_after = Pt(0)
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run(f"EMI STEM  |  {s['name']}  |  ")
    r.font.name, r.font.size = "Noto Sans", Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    return doc


def front(doc, s, kind):
    title = "Lesson Plan Template" if kind == "template" else "Lesson Plan Example"
    vietnamese = "Mẫu kế hoạch bài dạy" if kind == "template" else "Giáo án minh họa"
    paragraph(doc, f"{s['name']} {title}", "Title")
    paragraph(doc, f"{s['vi']}  |  {vietnamese}", "Subtitle")
    bilingual(doc, s["topic"], "Topic / Chủ đề")
    paragraph(doc, "English-Vietnamese EMI resource independently prepared using Cambridge's planning structure. / Tài liệu EMI Anh-Việt biên soạn độc lập theo cấu trúc lập kế hoạch của Cambridge.", "Reference")
    source = f"{s['book']}\n{s['chapter']}"
    paragraph(doc, source, "Reference")
    if s.get("syllabusBased"):
        paragraph(doc, "Engineering is based on the Design & Technology syllabus, not a supplied full textbook. / Engineering dựa trên syllabus Design & Technology, không phải sách giáo khoa đầy đủ đã cung cấp.", "Reference")


def phase_minutes(s):
    phases = {stage: sum(row[0] for row in s["stages"] if row[1] == stage)
              for stage in ["Beginning", "Middle", "End"]}
    assert sum(phases.values()) == LESSON_MINUTES
    return phases


def metadata(doc, s, example=False):
    rows = [
        ["School and teacher / Trường và giáo viên", "......................................................"],
        ["Date and class / Ngày và lớp", "......................................................"],
        ["Present and absent / Có mặt và vắng", "......................................................"],
        ["Level and reference / Trình độ và mã", f"IGCSE {s['code']}; ages 14-16; {LESSON_MINUTES} minutes\nĐộ tuổi 14-16; {LESSON_MINUTES} phút" if example else f"IGCSE {s['code']}; {LESSON_MINUTES} minutes / phút"]
    ]
    table(doc, ["Planning context / Bối cảnh", "Teacher details / Thông tin"], rows, [2.8, 4.4], 10)


def vocab(doc, s):
    table(doc, ["English term / Thuật ngữ", "Vietnamese meaning / Nghĩa tiếng Việt"], s["vocabulary"], [3.3, 3.9], 11)


def reflection(doc, compact=False):
    heading(doc, "Reflection and evaluation", "Phản tư và đánh giá sau bài dạy")
    paragraph(doc, "Complete after teaching using learner evidence; no outcomes are assumed. / Điền sau khi dạy dựa trên bằng chứng học tập; không giả định kết quả.", "Reference")
    blank(doc, "Evidence against objectives and criteria", "Bằng chứng theo mục tiêu và tiêu chí")
    blank(doc, "Support and timing adjustments that worked or need revision", "Hỗ trợ và thời gian hiệu quả hoặc cần điều chỉnh")
    blank(doc, "Next lesson action for the class or individual learners", "Việc cần làm ở bài sau cho lớp hoặc từng học sinh")
    if not compact:
        blank(doc, "One successful aspect and one improvement with reasons", "Một điểm tốt và một điểm cần cải thiện kèm lý do")


def references(doc, s, include_supplement=True):
    paragraph(doc, "Sources and authorship / Nguồn và biên soạn", "Heading 2")
    paragraph(doc, "Planning structure adapted from Cambridge guidance; EMI support, activities and questions are independently authored, not copied textbook exercises. Not issued or endorsed by Cambridge. Verify the syllabus for your examination year. / Cấu trúc dựa trên hướng dẫn Cambridge; hỗ trợ EMI, hoạt động và câu hỏi được soạn độc lập. Không do Cambridge phát hành hoặc chứng nhận. Kiểm tra syllabus của năm thi.", "Reference")
    link(doc, "Cambridge International lesson-planning guidance / Hướng dẫn lập kế hoạch bài dạy", COURSE)
    paragraph(doc, "Structural reference: Cambridge IGCSE Biology Teacher Guide, Appendix, p. 24. / Tham chiếu cấu trúc: phụ lục mẫu kế hoạch, trang 24.", "Reference")
    link(doc, "Referenced book or syllabus in Module 4 / Sách hoặc syllabus tham chiếu tại Module 4", f"{SITE}/api/books?subject={s['key']}&file={s['bookId']}")
    if include_supplement and s.get("supplement"):
        paragraph(doc, "Complementary publisher reference only / Sách tham khảo bổ sung: Collins, Cambridge IGCSE Design & Technology Student's Book, second edition, 2018. ISBN 9780008293277. Full book not supplied / Chưa có sách đầy đủ.", "Reference")
        link(doc, "Collins publisher overview / Giới thiệu sách của nhà xuất bản", s["supplementUrl"])


def template(s):
    d = new_document(s, "template")
    front(d, s, "template")
    metadata(d, s)
    bilingual(d, ["Adapt the topic and book reference to your lesson; complete both languages consistently.", "Điều chỉnh chủ đề và nguồn sách theo bài; điền hai ngôn ngữ thống nhất."])
    blank(d, "Curriculum learning aim", "Mục tiêu học tập theo chương trình")
    blank(d, "Lesson objectives using observable verbs up to three", "Mục tiêu bài dạy với động từ quan sát được tối đa ba")
    blank(d, "Success criteria and evidence of achievement", "Tiêu chí thành công và bằng chứng đạt mục tiêu")
    blank(d, "Language outcome for explaining subject reasoning", "Mục tiêu ngôn ngữ để giải thích lập luận môn học")
    heading(d, "Vocabulary and prior learning", "Từ vựng và kiến thức nền", new_page=True)
    paragraph(d, "Suggested topic vocabulary; revise for your chosen lesson. / Từ vựng gợi ý cho chủ đề; điều chỉnh theo bài đã chọn.")
    vocab(d, s)
    blank(d, "Sentence frame", "Mẫu câu")
    blank(d, "Prior knowledge and misconception check", "Kiểm tra kiến thức nền và hiểu nhầm")
    heading(d, "Plan", "Tiến trình bài dạy")
    phases = phase_minutes(s)
    table(d, ["Time / Thời gian", "Teacher and learner activity / Hoạt động", "Resources / Học liệu"], [
        [f"Beginning\nMở đầu\n{phases['Beginning']} min", "Prior-learning prompt and shared goals / Câu hỏi nền và mục tiêu\n...............................................................", "...................."],
        [f"Middle\nPhát triển\n{phases['Middle']} min", "Model, tasks and checks / Minh họa, luyện tập và kiểm tra\n...............................................................\n...............................................................", "....................\n...................."],
        [f"End\nKết thúc\n{phases['End']} min", "Exit evidence and next step / Bằng chứng cuối giờ và bước tiếp\n...............................................................", "...................."]
    ], [1, 4.9, 1.3], 10)
    paragraph(d, f"Total / Tổng: {LESSON_MINUTES} minutes / phút. Split the middle into timed steps as needed. / Chia phần phát triển thành các bước có thời gian khi cần.", "Reference")
    heading(d, "Additional planning", "Thông tin bổ sung", new_page=True)
    blank(d, "Support and access arrangements", "Hỗ trợ và điều chỉnh tiếp cận", 2)
    blank(d, "Extension that deepens reasoning", "Mở rộng để phát triển tư duy")
    blank(d, "Assessment checks and response to incomplete learning", "Kiểm tra đánh giá và xử lý khi chưa đạt", 2)
    blank(d, "Health safety and purposeful ICT use", "Sức khỏe an toàn và ứng dụng CNTT có mục đích")
    blank(d, "Homework linked to objectives and checking method", "Bài về nhà gắn mục tiêu và cách kiểm tra")
    reflection(d, compact=True)
    references(d, s, include_supplement=False)
    return d


def example(s):
    d = new_document(s, "example")
    front(d, s, "example")
    metadata(d, s, True)
    paragraph(d, "Illustrative context: mixed English proficiency; pairs and groups of three. Confirm readiness and adjust locally. / Bối cảnh minh họa: trình độ tiếng Anh khác nhau; làm cặp và nhóm ba. Kiểm tra mức sẵn sàng và điều chỉnh thực tế.", "Reference")
    bilingual(d, s["learning"], "Curriculum learning aim / Mục tiêu theo chương trình")
    table(d, ["Lesson objectives / Mục tiêu bài dạy", "Success criteria / Tiêu chí thành công"], [
        [f"{i}. {goal[0]}\n{goal[1]}", f"{criterion[0]}\n{criterion[1]}"]
        for i, (goal, criterion) in enumerate(zip(s["goals"], s["success"]), 1)
    ], [3.6, 3.6], 10.5)
    heading(d, "Vocabulary and prior learning", "Từ vựng và kiến thức nền", new_page=True)
    vocab(d, s)
    paragraph(d, "EMI language outcome / Mục tiêu ngôn ngữ EMI", "Heading 2")
    bilingual(d, ["Explain one subject decision in English using accurate terminology and a sentence frame. Vietnamese rehearsal is allowed; assess subject understanding independently of language accuracy.", "Giải thích một quyết định của môn học bằng tiếng Anh với thuật ngữ chính xác và mẫu câu. Cho phép chuẩn bị bằng tiếng Việt; đánh giá hiểu biết môn học độc lập với độ chính xác ngôn ngữ."])
    for pair in s["frames"]:
        bilingual(d, pair)
    bilingual(d, s["prior"], "Prior learning check / Kiểm tra kiến thức nền")
    bilingual(d, s["misconception"], "Anticipated misconception / Hiểu nhầm dự kiến")
    heading(d, "Plan", "Tiến trình bài dạy", new_page=True)
    phases = phase_minutes(s)
    paragraph(d, f"{LESSON_MINUTES} minutes: beginning {phases['Beginning']}, middle {phases['Middle']}, end {phases['End']}. / {LESSON_MINUTES} phút: mở đầu {phases['Beginning']}, phát triển {phases['Middle']}, kết thúc {phases['End']}.")
    rows, elapsed = [], 0
    labels = {"Beginning": "Mở đầu", "Middle": "Phát triển", "End": "Kết thúc"}
    for duration, stage, en, vi, resource in s["stages"]:
        rows.append([f"{elapsed}-{elapsed + duration}\n{stage}\n{labels[stage]}", en + "\n" + vi, resource])
        elapsed += duration
    assert elapsed == LESSON_MINUTES
    table(d, ["Minutes / Phút", "Planned teaching and learning / Dạy và học", "Resources / Học liệu"], rows, [.85, 4.95, 1.4], 10.5)
    bilingual(d, ["Use the learner worksheet section. Prepare copies and the board model before the lesson; collect individual exit evidence rather than assuming group success reflects every learner.", "Dùng phần phiếu học tập. Chuẩn bị bản in và mô hình trên bảng trước giờ dạy; thu bằng chứng cuối giờ cá nhân, không suy ra mọi học sinh đều đạt từ kết quả nhóm."])
    heading(d, "Additional planning", "Thông tin bổ sung", new_page=True)
    bilingual(d, s["support"], "Support / Hỗ trợ")
    bilingual(d, s["challenge"], "Challenge / Mở rộng")
    bilingual(d, s["assessment"], "Assessment and next steps / Đánh giá và bước tiếp")
    bilingual(d, s["safety"], "Health safety and ICT / An toàn và CNTT")
    bilingual(d, s["homework"], "Homework / Bài về nhà")
    heading(d, "Learner worksheet", "Phiếu học tập", new_page=True)
    paragraph(d, "Name / Họ tên: ...................................  Class / Lớp: ................", "Reference")
    for i, pair in enumerate(s["tasks"], 1):
        bilingual(d, [f"{i}. {pair[0]}", f"{i}. {pair[1]}"])
        if i == 1 and s.get("dataset"):
            table(d, s["dataset"][0], s["dataset"][1:], [.65, 1.65, 1.5, 1.6, 1.8], 9.5)
    paragraph(d, "Use your workbook or a separate sheet for full responses. / Trình bày lời giải đầy đủ trong vở hoặc trên giấy riêng.", "Reference")
    bilingual(d, s["exit"], "Individual exit question / Câu hỏi cuối giờ cá nhân")
    blank(d, "My answer and reasoning", "Câu trả lời và lập luận của em", 2)
    heading(d, "Answer guidance", "Hướng dẫn đáp án", new_page=True)
    for pair in s["answers"]:
        bilingual(d, pair)
    bilingual(d, s["exitAnswer"], "Exit answer / Đáp án cuối giờ")
    bilingual(d, s["homeworkAnswer"], "Homework answer / Đáp án bài về nhà")
    heading(d, "Reflection and evaluation", "Phản tư và đánh giá sau bài dạy", new_page=True)
    paragraph(d, "Complete after teaching using actual evidence. / Điền sau khi dạy dựa trên bằng chứng thực tế.", "Reference")
    blank(d, "Which objectives were achieved and what is the evidence", "Mục tiêu nào đạt và bằng chứng là gì", 2)
    blank(d, "Which learners need further support", "Học sinh nào cần hỗ trợ thêm", 2)
    blank(d, "How well did the language support work", "Hỗ trợ ngôn ngữ hiệu quả thế nào", 2)
    blank(d, "What timing or grouping changes were needed", "Đã cần đổi thời gian hoặc cách chia nhóm thế nào", 2)
    blank(d, "One strength and one improvement with reasons", "Một điểm tốt và một điểm cần cải thiện kèm lý do", 2)
    blank(d, "Action for the next lesson", "Việc cần làm ở bài sau", 2)
    references(d, s)
    return d


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--subject", choices=[s["key"] for s in DATA["subjects"]])
    args = parser.parse_args()
    OUT.mkdir(parents=True, exist_ok=True)
    for s in DATA["subjects"]:
        if args.subject and args.subject != s["key"]:
            continue
        for kind, build in [("template", template), ("example", example)]:
            destination = OUT / f"{s['key']}-{kind}-en-vi.docx"
            build(s).save(destination)
            print(destination.name)


if __name__ == "__main__":
    main()
